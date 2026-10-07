#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Genera MANUAL-EQUIPO.docx: la misma version del manual que el PDF, pero en
formato Word para poder subirla a Google Docs y editarla ahi.

El contenido NO se duplica: este script reutiliza el bloque de contenido de
generar-manual.py (lo que va entre los comentarios "contenido" y "armado") y
solo cambia como se dibuja cada cosa.

Uso: python3 docs/kit-nuevo-repo/generar-manual-docx.py
"""

import os
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import nsdecls, qn
from docx.oxml import parse_xml
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
FUENTE = os.path.join(HERE, "generar-manual.py")
SALIDA = os.path.join(HERE, "MANUAL-EQUIPO.docx")

VERDE = RGBColor(0x1D, 0x7A, 0x5F)
GRIS = RGBColor(0x3C, 0x46, 0x50)
TINTA = RGBColor(0x1F, 0x25, 0x2B)
ROJO = RGBColor(0xB0, 0x30, 0x30)
DORADO = RGBColor(0x8A, 0x6D, 0x1F)

VERDE_HEX = "1d7a5f"
VERDE_CLARO_HEX = "e6f4ef"
ROJO_CLARO_HEX = "fdecec"
DORADO_CLARO_HEX = "fdf6e3"
GRIS_CLARO_HEX = "f1f3f5"
CODIGO_HEX = "f5f7f9"

cm = 28.3464567  # 1 cm en puntos, para entender los anchos del script del PDF

doc = Document()

# ------------------------------------------------------------------ estilos
sec = doc.sections[0]
sec.left_margin = Cm(2)
sec.right_margin = Cm(2)
sec.top_margin = Cm(2)
sec.bottom_margin = Cm(2)
ANCHO_UTIL = sec.page_width - sec.left_margin - sec.right_margin

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.font.color.rgb = TINTA
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.15

for nombre, tam, color in (("Heading 1", 18, VERDE), ("Heading 2", 13, GRIS)):
    st = doc.styles[nombre]
    st.font.name = "Calibri"
    st.font.size = Pt(tam)
    st.font.bold = True
    st.font.color.rgb = color
    st.paragraph_format.space_before = Pt(16 if tam == 18 else 12)
    st.paragraph_format.space_after = Pt(8 if tam == 18 else 5)
    st.paragraph_format.keep_with_next = True


def _sombrear(celda, hexcolor):
    celda._tc.get_or_add_tcPr().append(
        parse_xml(r'<w:shd {} w:val="clear" w:color="auto" w:fill="{}"/>'.format(
            nsdecls("w"), hexcolor)))


def _borde(tabla, hexcolor="c9d1d9", tam=4):
    tbl_pr = tabla._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + lado)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(tam))
        el.set(qn("w:color"), hexcolor)
        borders.append(el)
    tbl_pr.append(borders)


# ------------------------------------------------ texto con etiquetas simples
TOKEN = re.compile(r"(<b>|</b>|<i>|</i>|<font[^>]*>|</font>|<br\s*/?>)")


def _escribir(par, texto, base=None):
    """Convierte el marcado sencillo del PDF (<b>, <i>, <font>, <br/>) a runs."""
    negrita = cursiva = mono = False
    for trozo in TOKEN.split(texto):
        if not trozo:
            continue
        if trozo == "<b>":
            negrita = True
        elif trozo == "</b>":
            negrita = False
        elif trozo == "<i>":
            cursiva = True
        elif trozo == "</i>":
            cursiva = False
        elif trozo.startswith("<font"):
            mono = True
        elif trozo == "</font>":
            mono = False
        elif re.match(r"<br\s*/?>", trozo):
            par.add_run().add_break()
        else:
            r = par.add_run(trozo.replace("&amp;", "&").replace("&lt;", "<")
                            .replace("&gt;", ">"))
            r.bold = negrita
            r.italic = cursiva
            if mono:
                r.font.name = "Consolas"
                r.font.size = Pt(10)
            if base:
                base(r)
    return par


# -------------------------------------------- mismos helpers que el PDF, en docx
def p(texto, estilo="p"):
    if estilo == "titulo":
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_after = Pt(6)
        _escribir(par, texto, lambda r: (setattr(r.font, "size", Pt(30)),
                                         setattr(r.font, "bold", True),
                                         setattr(r.font.color, "rgb", VERDE)))
    elif estilo == "subtitulo":
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _escribir(par, texto, lambda r: (setattr(r.font, "size", Pt(13)),
                                         setattr(r.font.color, "rgb", GRIS)))
    elif estilo == "h1":
        par = doc.add_paragraph(style="Heading 1")
        _escribir(par, texto)
    elif estilo == "h2":
        par = doc.add_paragraph(style="Heading 2")
        _escribir(par, texto)
    else:
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _escribir(par, texto)
    return par


def _numeracion_nueva():
    """Crea una numeracion propia para que cada lista vuelva a empezar en 1."""
    numbering = doc.part.numbering_part.element
    abstracto = doc.styles["List Number"].element.find(qn("w:pPr")).find(qn("w:numPr"))
    abs_id = None
    if abstracto is not None:
        num_id_ref = abstracto.find(qn("w:numId")).get(qn("w:val"))
        for num in numbering.findall(qn("w:num")):
            if num.get(qn("w:numId")) == num_id_ref:
                abs_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
    if abs_id is None:
        return None
    nuevo_id = max([int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(nuevo_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), abs_id)
    num.append(ref)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    inicio = OxmlElement("w:startOverride")
    inicio.set(qn("w:val"), "1")
    override.append(inicio)
    num.append(override)
    numbering.append(num)
    return nuevo_id


def vinetas(items, numerada=False):
    num_id = _numeracion_nueva() if numerada else None
    for t in items:
        par = doc.add_paragraph(style="List Number" if numerada else "List Bullet")
        par.paragraph_format.space_after = Pt(4)
        if num_id is not None:
            pPr = par._p.get_or_add_pPr()
            numPr = OxmlElement("w:numPr")
            ilvl = OxmlElement("w:ilvl")
            ilvl.set(qn("w:val"), "0")
            nid = OxmlElement("w:numId")
            nid.set(qn("w:val"), str(num_id))
            numPr.append(ilvl)
            numPr.append(nid)
            pPr.append(numPr)
        _escribir(par, t)
    return None


def codigo(texto):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _borde(t, "d6dce2")
    celda = t.cell(0, 0)
    celda.width = ANCHO_UTIL
    _sombrear(celda, CODIGO_HEX)
    lineas = texto.split("\n")
    for i, linea in enumerate(lineas):
        par = celda.paragraphs[0] if i == 0 else celda.add_paragraph()
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.line_spacing = 1
        r = par.add_run(linea)
        r.font.name = "Consolas"
        r.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return None


def caja(titulo, texto, fondo, borde_hex, color_titulo):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _borde(t, borde_hex, 8)
    celda = t.cell(0, 0)
    celda.width = ANCHO_UTIL
    _sombrear(celda, fondo)
    par = celda.paragraphs[0]
    par.paragraph_format.space_after = Pt(2)
    r = par.add_run(titulo)
    r.bold = True
    r.font.color.rgb = color_titulo
    par2 = celda.add_paragraph()
    par2.paragraph_format.space_after = Pt(0)
    par2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _escribir(par2, texto)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return None


def alerta(titulo, texto):
    return caja(titulo, texto, ROJO_CLARO_HEX, "b03030", ROJO)


def tip(titulo, texto):
    return caja(titulo, texto, VERDE_CLARO_HEX, "1d7a5f", VERDE)


def ejemplo(titulo, texto):
    return caja(titulo, texto, DORADO_CLARO_HEX, "8a6d1f", DORADO)


def tabla(encabezados, filas, anchos):
    t = doc.add_table(rows=1, cols=len(encabezados))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _borde(t)
    anchos_cm = [Cm(a / cm) for a in anchos]
    cabecera = t.rows[0].cells
    for i, h in enumerate(encabezados):
        cabecera[i].width = anchos_cm[i]
        _sombrear(cabecera[i], VERDE_HEX)
        par = cabecera[i].paragraphs[0]
        par.paragraph_format.space_after = Pt(0)
        r = par.add_run(re.sub(r"<[^>]+>", "", h))
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    # que la fila de titulos se repita si la tabla parte de pagina
    trPr = t.rows[0]._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:tblHeader"))
    for n, fila in enumerate(filas):
        celdas = t.add_row().cells
        for i, c in enumerate(fila):
            celdas[i].width = anchos_cm[i]
            if n % 2 == 1:
                _sombrear(celdas[i], GRIS_CLARO_HEX)
            par = celdas[i].paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            _escribir(par, c, lambda r: setattr(r.font, "size", Pt(10)))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return None


class PageBreak(object):
    def __init__(self):
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


class Spacer(object):
    def __init__(self, ancho, alto):
        if alto >= 0.8 * cm:
            par = doc.add_paragraph()
            par.paragraph_format.space_after = Pt(alto * 0.6)


# ---------------------------------------------- pie de pagina con numeracion
def _pie():
    par = sec.footer.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    par.add_run("MonteCristo · Manual del equipo\t\t")
    for texto, campo in (("PAGE", True),):
        pass
    r = par.add_run("Página ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    rin = OxmlElement("w:r")
    tin = OxmlElement("w:t")
    tin.text = "1"
    rin.append(tin)
    fld.append(rin)
    par._p.append(fld)
    for run in par.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x8A, 0x94, 0x9E)


# ------------------------------------------------------------------- armado
fuente = open(FUENTE, encoding="utf-8").read().split("\n")
ini = next(i for i, l in enumerate(fuente) if l.startswith("# ---") and "contenido" in l)
fin = next(i for i, l in enumerate(fuente) if l.startswith("# ---") and "armado" in l)
contenido = "\n".join(fuente[ini + 1:fin])

exec(compile(contenido, "contenido-del-manual", "exec"), globals())

_pie()
doc.core_properties.title = "MonteCristo - Manual del equipo"
doc.core_properties.author = "Equipo MonteCristo"
doc.save(SALIDA)

print("OK: %s (%.0f KB)" % (os.path.basename(SALIDA), os.path.getsize(SALIDA) / 1024))
